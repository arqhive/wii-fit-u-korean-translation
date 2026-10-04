"""번역 작업 분할/병합

  python tl_split.py split [목표글자수]   translation/work/in_NN.json 생성 (기본 15000자 단위)
  python tl_split.py merge               translation/work/out_*.json 검증 후 엑셀 E열에 반영
"""
import os, sys, json, glob, re, collections
from bmg_tool import load_orig, tokens, plain, XLSX

WORK = os.path.join(os.path.dirname(XLSX), 'work')
PUA = re.compile('[\ue000-\uf8ff]')
KANA_KANJI = re.compile('[぀-ヿ一-鿿ｦ-ﾟ]')
# 폰트 문자표 페이지(대사 아님) -> 번역 제외
EXCLUDE = set(range(6921, 7036)) | {8495}


def needs_translation(i, t):
    return i not in EXCLUDE and bool(KANA_KANJI.search(plain(t)))


def split():
    target = int(sys.argv[2]) if len(sys.argv) > 2 else 15000
    _, _, bmg = load_orig()
    os.makedirs(WORK, exist_ok=True)
    chunks = [[]]; size = 0; skipped = 0
    for i, t in enumerate(bmg.texts):
        if not needs_translation(i, t):
            skipped += 1
            continue
        if size >= target:
            chunks.append([]); size = 0
        chunks[-1].append({'i': i, 'jp': t})
        size += len(plain(t))
    for n, ch in enumerate(chunks):
        txt = json.dumps(ch, ensure_ascii=False, indent=0)
        txt = PUA.sub(lambda m: '\\u%04x' % ord(m.group()), txt)  # 아이콘 문자를 보이는 \uXXXX 이스케이프로
        with open(os.path.join(WORK, f'in_{n:02d}.json'), 'w', encoding='utf-8') as f:
            f.write(txt)
        print(f'in_{n:02d}.json: {len(ch)}문장 #{ch[0]["i"]}~#{ch[-1]["i"]}, {sum(len(plain(x["jp"])) for x in ch)}자')
    print(f'번역 대상 아님(기호·영문·빈 문장·문자표): {skipped}개')


def check(jp, ko):
    errs = []
    tj = sorted(v for k, v in tokens(jp) if k == 'tag')
    try:
        tk = sorted(v for k, v in tokens(ko) if k == 'tag')
    except ValueError as e:
        return [f'중괄호 오류: {e}']
    if tj != tk:
        errs.append(f'태그 불일치 누락{dict(collections.Counter(tj) - collections.Counter(tk))} 추가{dict(collections.Counter(tk) - collections.Counter(tj))}')
    if sorted(PUA.findall(jp)) != sorted(PUA.findall(ko)):
        errs.append(f'아이콘 문자 불일치 {[hex(ord(c)) for c in PUA.findall(jp)]} -> {[hex(ord(c)) for c in PUA.findall(ko)]}')
    if re.search('[぀-ヿ]', plain(ko)):
        errs.append('가나 잔존')
    lj, lk = jp.count('\n'), ko.count('\n')
    if abs(lj - lk) > 1:
        errs.append(f'줄 수 차이 {lj + 1}->{lk + 1}')
    return errs


def merge():
    from openpyxl import load_workbook
    _, _, bmg = load_orig()
    got = {}
    for p in sorted(glob.glob(os.path.join(WORK, 'out_*.json'))):
        with open(p, encoding='utf-8') as f:
            for k, v in json.load(f).items():
                got[int(k)] = v
    expected = [i for i, t in enumerate(bmg.texts) if needs_translation(i, t)]
    missing = [i for i in expected if i not in got]
    problems = {}
    for i, ko in got.items():
        e = check(bmg.texts[i], ko)
        if e:
            problems[i] = e
    with open(os.path.join(WORK, 'merge_report.txt'), 'w', encoding='utf-8') as f:
        f.write(f'번역 {len(got)} / 대상 {len(expected)}, 누락 {len(missing)}, 문제 {len(problems)}\n')
        f.write('누락: ' + ','.join(map(str, missing)) + '\n')
        for i, e in sorted(problems.items()):
            f.write(f'#{i}: {"; ".join(e)}\n   JP: {bmg.texts[i]!r}\n   KO: {got[i]!r}\n')
    wb = load_workbook(XLSX)
    ws = wb['messages']
    for row in ws.iter_rows(min_row=2):
        idx = row[0].value
        if idx in got:
            row[4].value = got[idx]
            row[5].value = ('검토 필요: ' + '; '.join(problems[idx])) if idx in problems else None
    wb.save(XLSX)
    print(f'번역 {len(got)} / 대상 {len(expected)}, 누락 {len(missing)}, 문제 {len(problems)} -> 엑셀 반영, 보고서 work/merge_report.txt')


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else ''
    {'split': split, 'merge': merge}.get(cmd, lambda: print(__doc__))()
