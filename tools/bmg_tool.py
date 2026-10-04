"""health_message.bmg <-> 번역 엑셀

사용법:
  python bmg_tool.py extract            원문 -> translation/health_message.xlsx
  python bmg_tool.py build              엑셀 번역 -> build/content/JP/Message/message.carc
  python bmg_tool.py roundtrip          수정 없이 재빌드해서 원본과 바이트 일치 확인
태그 표기: 제어코드 0x001A 는 {hex} 로 표시 (예: {ff00000100} 색상, {6400010006} 대기, {000000} 변수)
글자 그대로의 중괄호는 {{ / }} 로 표기
"""
import os, re, struct, sys
from wfu_lib import *

REL = r'JP\Message\message.carc'
XLSX = os.path.join(ROOT, 'translation', 'health_message.xlsx')
OUT = os.path.join(BUILD, 'content', 'JP', 'Message', 'message.carc')
TOK_RE = re.compile(r'\{\{|\}\}|\{([0-9a-f]+)\}|[{}]')


def tokens(text):
    """-> [(kind, value)] kind: 'text' | 'tag'"""
    out = []; pos = 0
    for m in TOK_RE.finditer(text):
        if m.start() > pos:
            out.append(('text', text[pos:m.start()]))
        s = m.group(0)
        if s == '{{':
            out.append(('text', '{'))
        elif s == '}}':
            out.append(('text', '}'))
        elif m.group(1):
            out.append(('tag', m.group(1)))
        else:
            raise ValueError(f'짝이 맞지 않는 중괄호 (글자 그대로는 {{{{ }}}} 사용): {text[:60]!r}')
        pos = m.end()
    if pos < len(text):
        out.append(('text', text[pos:]))
    return out


def plain(text):
    return ''.join(v for k, v in tokens(text) if k == 'text')


class BMG:
    def __init__(self, b):
        assert b[:8] == b'MESGbmg1'
        self.raw = b
        self.encoding = b[0x10]
        assert self.encoding == 2, 'UTF-16 BMG만 지원'
        self.secs = []; off = 0x20
        nsec = struct.unpack('>I', b[0x0C:0x10])[0]
        for _ in range(nsec):
            mg = b[off:off + 4]; sz = struct.unpack('>I', b[off + 4:off + 8])[0]
            self.secs.append((mg, off, sz)); off += sz
        sec = {m: (o, s) for m, o, s in self.secs}
        io, isz = sec[b'INF1']; self.dat_off, self.dat_sz = sec[b'DAT1']
        self.count, self.esz = struct.unpack('>HH', b[io + 8:io + 12])
        self.inf_off = io
        mo, _ = sec[b'MID1']
        self.mids = [struct.unpack('>I', b[mo + 16 + i * 4:mo + 20 + i * 4])[0] for i in range(self.count)]
        self.ptrs = []; self.attrs = []
        for i in range(self.count):
            e = io + 0x10 + i * self.esz
            self.ptrs.append(struct.unpack('>I', b[e:e + 4])[0])
            self.attrs.append(b[e + 4:e + self.esz])
        self.texts = [self._decode(p) for p in self.ptrs]

    def _decode(self, p):
        b = self.raw; q = self.dat_off + 8 + p; out = []
        while True:
            c = struct.unpack('>H', b[q:q + 2])[0]
            if c == 0:
                break
            if c == 0x1A:
                ln = b[q + 2]
                out.append('{' + b[q + 3:q + ln].hex() + '}'); q += ln; continue
            ch = chr(c)
            out.append(ch * 2 if ch in '{}' else ch); q += 2
        return ''.join(out)

    @staticmethod
    def encode(text):
        out = bytearray()
        for kind, v in tokens(text):
            if kind == 'text':
                out += v.encode('utf-16-be')
            else:
                rest = bytes.fromhex(v)
                out += b'\x00\x1a' + bytes([3 + len(rest)]) + rest
        return bytes(out + b'\0\0')

    def build(self, texts):
        b = self.raw
        # 원본 DAT1 저장 순서(오프셋 순) 유지
        order = sorted(range(self.count), key=lambda i: self.ptrs[i])
        first_gap = min(self.ptrs)  # DAT1 선두 여유 바이트(보통 빈 문자열 등)
        data = bytearray(b[self.dat_off + 8:self.dat_off + 8 + first_gap])
        newptr = [0] * self.count
        for i in order:
            newptr[i] = len(data)
            data += self.encode(texts[i])
        data += b'\0' * ((-len(data)) % 32)  # DAT1 본문(헤더 8바이트 제외)은 32바이트 정렬
        inf = bytearray(b[self.inf_off:self.inf_off + struct.unpack('>I', b[self.inf_off + 4:self.inf_off + 8])[0]])
        for i in range(self.count):
            struct.pack_into('>I', inf, 0x10 + i * self.esz, newptr[i])
        out = bytearray(b[:0x20])
        for mg, off, sz in self.secs:
            if mg == b'INF1':
                out += inf
            elif mg == b'DAT1':
                out += b'DAT1' + struct.pack('>I', len(data) + 8) + data
            else:
                sec = b[off:off + sz]
                out += sec + b'\0' * (sz - len(sec))  # 원본은 마지막 섹션 끝 8바이트가 잘려 있음
        delta = len(out) - (self.secs[-1][1] + self.secs[-1][2])
        struct.pack_into('>I', out, 8, struct.unpack('>I', b[8:12])[0] + delta)
        body_len = len(b) - (self.secs[-1][1] + self.secs[-1][2])  # 음수면 원본이 섹션 크기보다 짧음
        if body_len < 0:
            out = out[:len(out) + body_len]
        else:
            out += b[len(b) - body_len:] if body_len else b''
        return bytes(out)


def load_orig():
    carc = open(src_content(REL), 'rb').read()
    files = u8_read(yaz0_decompress(carc))
    assert len(files) == 1 and files[0][0] == 'health_message.bmg'
    return carc, files, BMG(files[0][1])


def extract():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    _, _, bmg = load_orig()
    wb = Workbook(); ws = wb.active; ws.title = 'messages'
    head = ['번호', 'MID', '속성', '원문(일본어)', '번역(한국어)', '비고', '원문 글자수']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='4472C4')
    for i in range(bmg.count):
        t = bmg.texts[i]
        ws.append([i, f'{bmg.mids[i]:08X}', bmg.attrs[i].hex(), t, None, None, len(plain(t))])
    for col, w in zip('ABCDEFG', (8, 11, 18, 60, 60, 20, 10)):
        ws.column_dimensions[col].width = w
    wrap = Alignment(wrap_text=True, vertical='top')
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.alignment = wrap
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = f'A1:G{bmg.count + 1}'
    g = wb.create_sheet('태그 설명')
    for r in [
        ['태그', '의미(추정)', '번역 시 주의'],
        ['{ff0000XXXX} ... {ff0000ffff}', '글자색 시작 / 원래 색으로', '감싸는 단어 위치에 맞게 옮기기'],
        ['{ff0001XXXX} ... {ff0001ffff}', '글자 크기/스타일 변경 / 복귀', '짝 유지'],
        ['{640001XXXX}', '대기(음성 싱크, 값=프레임)', '문장 흐름에 맞는 위치에 두기, 개수 유지 권장'],
        ['{640002XXXX}', '표정/애니메이션 신호(추정)', '문장 앞부분 위치 유지'],
        ['{0000XX}', '변수 삽입(이름·숫자 등)', '반드시 유지, 조사(은/는, 이/가) 주의'],
        ['{020001XXXX}', '선택지/입력 대기(추정)', '위치 유지'],
        ['줄바꿈', '셀 안 줄바꿈 = 게임 줄바꿈', '원문 줄 수를 크게 넘기지 않기'],
        ['{{ / }}', '글자 그대로의 중괄호 { }', '태그와 구분하기 위해 두 번 적기'],
    ]:
        g.append(r)
    for col, w in zip('ABC', (32, 34, 44)):
        g.column_dimensions[col].width = w
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    if os.path.exists(XLSX):
        sys.exit(f'이미 존재함(덮어쓰지 않음): {XLSX}')
    wb.save(XLSX)
    print(f'추출 완료: {bmg.count}개 -> {XLSX}')


def tags_of(t):
    return sorted(v for k, v in tokens(t) if k == 'tag')


def build():
    from openpyxl import load_workbook
    carc, files, bmg = load_orig()
    texts = list(bmg.texts); used = 0; warn = 0
    if '--idtest' in sys.argv:
        # 진단용: 모든 문장을 "번호 한글" 로 교체(태그는 순서대로 유지) -> 화면에 뜬 문장 번호 확인용
        for i, t in enumerate(bmg.texts):
            if not plain(t).strip():
                continue
            tags = ''.join('{' + v + '}' for k, v in tokens(t) if k == 'tag')
            texts[i] = f'{i}한글' + tags
            used += 1
        rows = []
    else:
        rows = load_workbook(XLSX, read_only=True)['messages'].iter_rows(min_row=2, values_only=True)
    for row in rows:
        idx, ko = row[0], row[4]
        if ko is None or str(ko).strip() == '':
            continue
        ko = str(ko).replace('\r\n', '\n').replace('_x000D_', '')
        if tags_of(ko) != tags_of(bmg.texts[idx]):
            warn += 1
            print(f'[태그 불일치] #{idx}: 원문 {tags_of(bmg.texts[idx])} / 번역 {tags_of(ko)}')
        texts[idx] = ko; used += 1
    new_bmg = bmg.build(texts)
    new_carc = yaz0_compress(u8_write([('health_message.bmg', new_bmg)]))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    open(OUT, 'wb').write(new_carc)
    # 폰트에 추가할 문자 후보: 실제로 바뀐(번역된) 문장의 문자만 (원문 그대로인 문장은 원래 폰트로 충분)
    chars = sorted({ch for i in range(bmg.count) if texts[i] != bmg.texts[i] for ch in plain(texts[i])})
    open(os.path.join(BUILD, 'used_chars.txt'), 'w', encoding='utf-8').write(''.join(chars))
    print(f'빌드 완료: 번역 적용 {used}개, 태그 경고 {warn}개 -> {OUT} ({len(new_carc)} bytes)')
    print(f'사용 문자 {len(chars)}종 (한글 {sum(1 for c in chars if 0xAC00 <= ord(c) <= 0xD7A3)}자) -> build/used_chars.txt')


def roundtrip():
    carc, files, bmg = load_orig()
    orig_bmg = files[0][1]
    print(f'메시지 {bmg.count}개, DAT1 오프셋 순서==번호 순서: {bmg.ptrs == sorted(bmg.ptrs)}, 최소 오프셋 {min(bmg.ptrs)}')
    re_bmg = bmg.build(bmg.texts)
    print('BMG 재빌드 일치:', re_bmg == orig_bmg, len(re_bmg), len(orig_bmg))
    if re_bmg != orig_bmg:
        k = next((i for i in range(min(len(re_bmg), len(orig_bmg))) if re_bmg[i] != orig_bmg[i]), None)
        print('  첫 차이 오프셋:', k)
    dec = yaz0_decompress(carc)
    re_u8 = u8_write(files)
    print('U8 재빌드 일치:', re_u8 == dec, len(re_u8), len(dec))
    comp = yaz0_compress(dec)
    print('Yaz0 압축->해제 일치:', yaz0_decompress(comp) == dec, f'압축 크기 {len(comp)} (원본 {len(carc)})')


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else ''
    {'extract': extract, 'build': build, 'roundtrip': roundtrip}.get(cmd, lambda: print(__doc__))()
